"""使用 python-docx 的 Word (.docx) extractor。"""

from __future__ import annotations

from pathlib import Path

from ..models import DocumentRecord, compute_doc_id
from ..utils import is_meaningful_text, normalize_text
from .base import BaseExtractor, ExtractionResult


_HEADING_PREFIXES = (
    "Heading 1",
    "Heading 2",
    "Heading 3",
    "Heading 4",
    "Heading 5",
    "Heading 6",
    "Title",
)


def _is_heading_style(style_name: str) -> bool:
    """判断 ``style_name`` 是否属于 Heading/Title 系列样式。"""
    return any(style_name.startswith(prefix) for prefix in _HEADING_PREFIXES)


def _format_table_as_text(rows: list[list[str]]) -> str:
    """把二维单元格字符串列表渲染为 Markdown 风格的表格。"""
    if not rows:
        return ""
    # 去掉完全为空的行。
    rows = [r for r in rows if any(cell.strip() for cell in r)]
    if not rows:
        return ""
    widths = [
        max(len((r[i] if i < len(r) else "").strip()) for r in rows)
        for i in range(max(len(r) for r in rows))
    ]
    lines: list[str] = []
    for r_idx, row in enumerate(rows):
        cells = [(row[i] if i < len(row) else "").strip().ljust(widths[i]) for i in range(len(widths))]
        lines.append("| " + " | ".join(cells) + " |")
        if r_idx == 0:
            lines.append("|" + "|".join("-" * (w + 2) for w in widths) + "|")
    return "\n".join(lines)


class WordExtractor(BaseExtractor):
    """基于 python-docx 的 Word (.docx) 文件抽取器。

    按 body 中段落 / 表格出现的顺序输出 record，标题（Heading/Title 样式）
    会作为单独 record 并填充 ``section``。
    """

    source_type: str = "docx"
    suffixes: tuple[str, ...] = (".docx",)

    def extract(
        self,
        path: Path,
        *,
        relative_path: str,
        file_sha256: str,
    ) -> ExtractionResult:
        try:
            from docx import Document  # type: ignore[import-not-found]
        except ImportError as exc:  # pragma: no cover - import guard
            raise RuntimeError(
                "python-docx is required for Word extraction. "
                "Install it via `pip install python-docx`."
            ) from exc

        result = ExtractionResult()
        current_section: str | None = None
        record_counter = 0

        document = Document(str(path))

        # 按顺序遍历文档 body，使段落和表格保持原本的交错顺序。
        body = document.element.body
        para_index = 0
        table_index = 0
        for child in body.iterchildren():
            tag = child.tag.split("}")[-1]
            if tag == "p":
                if para_index >= len(document.paragraphs):
                    continue
                paragraph = document.paragraphs[para_index]
                para_index += 1
                text = normalize_text(paragraph.text)
                style_name = paragraph.style.name if paragraph.style else ""
                if _is_heading_style(style_name) and is_meaningful_text(text):
                    current_section = text
                    # 把标题本身也作为一条 record 输出，避免丢失。
                    record_counter += 1
                    record = DocumentRecord(
                        doc_id=compute_doc_id(
                            relative_path,
                            record_index=record_counter,
                            text=text,
                            section=current_section,
                        ),
                        source_type=self.source_type,
                        source_path=relative_path,
                        file_name=path.name,
                        section=current_section,
                        record_index=record_counter,
                        text=text,
                        file_sha256=file_sha256,
                        metadata={"style": style_name, "kind": "heading"},
                    )
                    result.records.append(record)
                    continue
                if not is_meaningful_text(text):
                    continue
                record_counter += 1
                record = DocumentRecord(
                    doc_id=compute_doc_id(
                        relative_path,
                        record_index=record_counter,
                        text=text,
                        section=current_section,
                    ),
                    source_type=self.source_type,
                    source_path=relative_path,
                    file_name=path.name,
                    section=current_section,
                    record_index=record_counter,
                    text=text,
                    file_sha256=file_sha256,
                    metadata={"style": style_name, "kind": "paragraph"},
                )
                result.records.append(record)
            elif tag == "tbl":
                if table_index >= len(document.tables):
                    continue
                table = document.tables[table_index]
                table_index += 1
                rows: list[list[str]] = []
                for row in table.rows:
                    rows.append([cell.text for cell in row.cells])
                if not any(any(cell.strip() for cell in r) for r in rows):
                    continue
                rendered = _format_table_as_text(rows)
                rendered = normalize_text(rendered)
                if not is_meaningful_text(rendered):
                    continue
                record_counter += 1
                record = DocumentRecord(
                    doc_id=compute_doc_id(
                        relative_path,
                        record_index=record_counter,
                        text=rendered,
                        section=current_section,
                    ),
                    source_type=self.source_type,
                    source_path=relative_path,
                    file_name=path.name,
                    section=current_section,
                    record_index=record_counter,
                    text=rendered,
                    file_sha256=file_sha256,
                    metadata={"kind": "table", "rows": len(rows)},
                )
                result.records.append(record)

        return result