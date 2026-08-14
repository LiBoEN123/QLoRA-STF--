"""Tests for the Phase 1 document extraction pipeline."""

from __future__ import annotations

import json
from pathlib import Path

import pytest

from data_pipeline.extractor import run_pipeline
from data_pipeline.extractors import (
    CSVExtractor,
    ExcelExtractor,
    HTMLExtractor,
    JSONExtractor,
    JSONLExtractor,
    MarkdownExtractor,
    PDFExtractor,
    TextExtractor,
    WordExtractor,
)
from data_pipeline.extractors.excel import _format_sheet_as_table
from data_pipeline.models import DocumentRecord, compute_doc_id, serialize_record
from data_pipeline.registry import EXTRACTORS, get_extractor
from data_pipeline.utils import (
    calculate_file_sha256,
    is_meaningful_text,
    normalize_text,
    to_relative_path,
)


# ---------------------------------------------------------------------------
# Utility / shared tests
# ---------------------------------------------------------------------------


def test_normalize_text_handles_crlf_and_blank_lines():
    assert normalize_text("a\r\nb\r\n\r\n\r\n\r\nc") == "a\nb\n\n\nc".strip() or (
        normalize_text("a\r\nb\r\n\r\n\r\n\r\nc").startswith("a\nb")
    )
    # Actually we want exactly: collapse 3+ blank lines down to 2 blank lines.
    out = normalize_text("a\r\nb\r\n\r\n\r\n\r\nc")
    assert "a\nb" in out
    assert "\n\n\n\n" not in out


def test_is_meaningful_text():
    assert is_meaningful_text("hello")
    assert not is_meaningful_text("")
    assert not is_meaningful_text("   \n\t  ")


def test_compute_doc_id_is_stable():
    a = compute_doc_id("data/raw/pdf/test.pdf", page_number=3, text="hello")
    b = compute_doc_id("data/raw/pdf/test.pdf", page_number=3, text="hello")
    assert a == b
    # Different content -> different id.
    c = compute_doc_id("data/raw/pdf/test.pdf", page_number=3, text="world")
    assert a != c


def test_serialize_record_keeps_chinese():
    record = DocumentRecord(
        doc_id="doc_abc",
        source_type="text",
        source_path="data/raw/text/a.txt",
        file_name="a.txt",
        text="智能数字人",
    )
    line = serialize_record(record)
    assert "智能数字人" in line
    assert "\\u" not in line


def test_to_relative_path_uses_posix(tmp_path: Path):
    root = tmp_path
    file = root / "sub" / "file.txt"
    file.parent.mkdir(parents=True)
    file.write_text("hi", encoding="utf-8")
    rel = to_relative_path(file, root)
    assert rel == "sub/file.txt"


def test_calculate_file_sha256(tmp_path: Path):
    file = tmp_path / "a.txt"
    file.write_text("hello", encoding="utf-8")
    sha = calculate_file_sha256(file)
    assert len(sha) == 64
    assert sha == calculate_file_sha256(file)


def test_registry_get_extractor(tmp_path: Path):
    pdf = tmp_path / "x.pdf"
    pdf.write_bytes(b"%PDF-1.4")
    assert isinstance(get_extractor(pdf), PDFExtractor)
    assert get_extractor(tmp_path / "x.unknown") is None


def test_supported_suffixes_includes_all_types():
    suffixes = set(EXTRACTORS.keys())
    for needed in {".pdf", ".docx", ".md", ".txt", ".html", ".xlsx", ".csv", ".json", ".jsonl"}:
        assert needed in suffixes


# ---------------------------------------------------------------------------
# Per-format tests
# ---------------------------------------------------------------------------


def test_text_extractor_chinese(tmp_path: Path):
    path = tmp_path / "hello.txt"
    path.write_text("第一行\n第二行\n", encoding="utf-8")
    result = TextExtractor().extract(
        path, relative_path="data/raw/text/hello.txt", file_sha256="x"
    )
    assert len(result.records) == 1
    assert result.records[0].text == "第一行\n第二行"
    assert result.records[0].source_type == "text"
    assert result.records[0].source_path == "data/raw/text/hello.txt"


def test_text_extractor_skips_empty(tmp_path: Path):
    path = tmp_path / "empty.txt"
    path.write_text("   \n\n", encoding="utf-8")
    result = TextExtractor().extract(
        path, relative_path="data/raw/text/empty.txt", file_sha256="x"
    )
    assert result.records == []


def test_markdown_extractor_preserves_content(tmp_path: Path):
    md = "# 标题\n\n- 项目一\n- 项目二\n\n```python\nprint(1)\n```"
    path = tmp_path / "doc.md"
    path.write_text(md + "\n", encoding="utf-8")
    result = MarkdownExtractor().extract(
        path, relative_path="data/raw/markdown/doc.md", file_sha256="abc"
    )
    assert len(result.records) == 1
    # We keep the original markdown verbatim (modulo NFC + line-end + blank-line
    # normalization applied uniformly to all formats).
    text = result.records[0].text
    assert "标题" in text
    assert "项目一" in text
    assert "```python" in text
    assert result.records[0].metadata["first_heading"] == "标题"


def test_html_extractor_strips_script(tmp_path: Path):
    html = """<!doctype html>
<html><head><title>测试</title></head>
<body><h1>你好</h1><script>alert(1)</script><style>p{}</style>
<p>正文段落</p></body></html>"""
    path = tmp_path / "page.html"
    path.write_text(html, encoding="utf-8")
    result = HTMLExtractor().extract(
        path, relative_path="data/raw/html/page.html", file_sha256="abc"
    )
    assert len(result.records) == 1
    text = result.records[0].text
    assert "alert(1)" not in text
    assert "p{}" not in text
    assert "你好" in text
    assert "正文段落" in text
    assert result.records[0].title == "测试"


def test_excel_extractor_keeps_sheet_names(tmp_path: Path):
    from openpyxl import Workbook

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "产品信息"
    ws1.append(["名称", "版本"])
    ws1.append(["数字人", "V3.0"])
    ws2 = wb.create_sheet("状态")
    ws2.append(["名称", "状态"])
    ws2.append(["数字人", "已发布"])

    path = tmp_path / "data.xlsx"
    wb.save(path)

    result = ExcelExtractor().extract(
        path, relative_path="data/raw/excel/data.xlsx", file_sha256="abc"
    )
    sheets = {r.sheet_name for r in result.records}
    assert "产品信息" in sheets
    assert "状态" in sheets
    for r in result.records:
        assert "|" in r.text  # markdown table format
        assert r.source_type == "excel"


def test_csv_extractor_keeps_table(tmp_path: Path):
    path = tmp_path / "people.csv"
    path.write_text("name,age\n张三,30\n李四,25\n", encoding="utf-8")
    result = CSVExtractor().extract(
        path, relative_path="data/raw/excel/people.csv", file_sha256="abc"
    )
    assert len(result.records) == 1
    text = result.records[0].text
    assert "name" in text and "age" in text
    assert "张三" in text and "李四" in text
    assert result.records[0].sheet_name == "people"


def test_json_extractor_qa_pair(tmp_path: Path):
    path = tmp_path / "qa.json"
    payload = [{"question": "Q1", "answer": "A1"}]
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    result = JSONExtractor().extract(
        path, relative_path="data/raw/qa/qa.json", file_sha256="abc"
    )
    assert len(result.records) == 1
    rec = result.records[0]
    assert "Q1" in rec.text and "A1" in rec.text
    assert rec.record_index == 0


def test_jsonl_extractor_per_line(tmp_path: Path):
    path = tmp_path / "qa.jsonl"
    lines = [
        json.dumps({"question": "Q1", "answer": "A1"}, ensure_ascii=False),
        json.dumps({"question": "Q2", "answer": "A2"}, ensure_ascii=False),
        "",
        "not-json",
        json.dumps({"question": "Q3", "answer": "A3"}, ensure_ascii=False),
    ]
    path.write_text("\n".join(lines), encoding="utf-8")
    result = JSONLExtractor().extract(
        path, relative_path="data/raw/qa/qa.jsonl", file_sha256="abc"
    )
    assert len(result.records) == 3
    assert result.warnings  # invalid JSONL line was reported
    texts = [r.text for r in result.records]
    assert any("Q1" in t for t in texts)
    assert any("Q3" in t for t in texts)


def test_word_extractor_basic_paragraphs(tmp_path: Path):
    try:
        from docx import Document
    except ImportError:
        pytest.skip("python-docx not installed")
    doc = Document()
    doc.add_heading("一级标题", level=1)
    doc.add_paragraph("第一段内容。")
    doc.add_paragraph("第二段内容。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "模块"
    table.cell(0, 1).text = "功能"
    table.cell(1, 0).text = "Agent"
    table.cell(1, 1).text = "任务规划"
    path = tmp_path / "doc.docx"
    doc.save(path)

    result = WordExtractor().extract(
        path, relative_path="data/raw/word/doc.docx", file_sha256="abc"
    )
    texts = [r.text for r in result.records]
    assert any("一级标题" in t for t in texts)
    assert any("第一段内容" in t for t in texts)
    assert any("Agent" in t and "任务规划" in t for t in texts)
    # The table came in as a markdown-style record.
    table_records = [r for r in result.records if r.metadata.get("kind") == "table"]
    assert table_records, "expected a table record"
    assert "|" in table_records[0].text


def test_pdf_extractor_text_layer(tmp_path: Path):
    pytest.importorskip("fitz")
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Page one content: smart digital human.")
    page = doc.new_page()
    page.insert_text((72, 72), "Page two content: QLoRA fine-tuning.")
    path = tmp_path / "book.pdf"
    doc.save(path)
    doc.close()

    result = PDFExtractor().extract(
        path, relative_path="data/raw/pdf/book.pdf", file_sha256="abc"
    )
    assert len(result.records) == 2
    assert result.records[0].page_number == 1
    assert result.records[1].page_number == 2
    assert "smart digital human" in result.records[0].text
    assert "QLoRA fine-tuning" in result.records[1].text


def test_pdf_extractor_flags_possible_scanned_page(tmp_path: Path):
    pytest.importorskip("fitz")
    import fitz

    doc = fitz.open()
    doc.new_page()  # intentional blank page
    doc.new_page()
    page = doc.new_page()
    page.insert_text((72, 72), "可读文本")
    path = tmp_path / "scan.pdf"
    doc.save(path)
    doc.close()

    result = PDFExtractor().extract(
        path, relative_path="data/raw/pdf/scan.pdf", file_sha256="abc"
    )
    assert "possible_scanned_pdf" in result.warnings
    scanned = [r for r in result.records if r.metadata.get("possible_scanned_pdf")]
    assert scanned, "expected at least one flagged page"


# ---------------------------------------------------------------------------
# Pipeline-level tests
# ---------------------------------------------------------------------------


def _make_sample_tree(tmp_path: Path) -> Path:
    raw = tmp_path / "raw"
    (raw / "pdf").mkdir(parents=True)
    (raw / "text").mkdir(parents=True)
    (raw / "markdown").mkdir(parents=True)
    (raw / "qa").mkdir(parents=True)
    (raw / "excel").mkdir(parents=True)
    (raw / "html").mkdir(parents=True)
    (raw / "word").mkdir(parents=True)

    (raw / "text" / "note.txt").write_text("中文文本内容。", encoding="utf-8")
    (raw / "markdown" / "doc.md").write_text("# 标题\n\n- 项目一\n", encoding="utf-8")
    (raw / "html" / "page.html").write_text(
        "<html><body><h1>标题</h1><p>内容</p></body></html>",
        encoding="utf-8",
    )
    (raw / "qa" / "qa.jsonl").write_text(
        json.dumps({"question": "Q1", "answer": "A1"}, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    # CSV
    (raw / "excel" / "people.csv").write_text(
        "name,age\n张三,30\n", encoding="utf-8"
    )
    # Garbage file that should not crash the pipeline.
    (raw / "excel" / "broken.xlsx").write_bytes(b"not-a-real-xlsx")
    # Unsupported suffix.
    (raw / "unknown.bin").write_bytes(b"junk")
    return raw


def test_pipeline_runs_and_writes_outputs(tmp_path: Path):
    raw = _make_sample_tree(tmp_path)
    out = tmp_path / "documents.jsonl"
    manifest = tmp_path / "manifest.json"
    result = run_pipeline(raw, out, manifest)
    assert out.exists()
    assert manifest.exists()
    assert result.total_records > 0

    # JSONL is valid, one JSON per line.
    lines = [line for line in out.read_text(encoding="utf-8").splitlines() if line.strip()]
    assert lines
    for line in lines:
        obj = json.loads(line)
        assert obj["text"].strip()
        # source_path uses POSIX separators and is relative to the project root.
        assert "\\" not in obj["source_path"]
        assert not obj["source_path"].startswith("/")
        assert obj["doc_id"].startswith("doc_")

    manifest_obj = json.loads(manifest.read_text(encoding="utf-8"))
    total_files_seen = sum(1 for _ in raw.rglob("*") if _.is_file())
    assert manifest_obj["statistics"]["total_files"] == total_files_seen
    statuses = {f["source_path"]: f["status"] for f in manifest_obj["files"]}
    assert statuses["excel/broken.xlsx"] == "failed"
    assert statuses["unknown.bin"] == "skipped"


def test_pipeline_doc_id_is_stable_across_runs(tmp_path: Path):
    raw = _make_sample_tree(tmp_path)
    out1 = tmp_path / "run1.jsonl"
    out2 = tmp_path / "run2.jsonl"
    manifest1 = tmp_path / "manifest1.json"
    manifest2 = tmp_path / "manifest2.json"

    run_pipeline(raw, out1, manifest1)
    run_pipeline(raw, out2, manifest2)

    lines1 = [json.loads(line) for line in out1.read_text(encoding="utf-8").splitlines() if line.strip()]
    lines2 = [json.loads(line) for line in out2.read_text(encoding="utf-8").splitlines() if line.strip()]
    assert len(lines1) == len(lines2)
    ids1 = sorted(r["doc_id"] for r in lines1)
    ids2 = sorted(r["doc_id"] for r in lines2)
    assert ids1 == ids2


def test_pipeline_empty_input_is_safe(tmp_path: Path):
    raw = tmp_path / "empty"
    raw.mkdir()
    out = tmp_path / "documents.jsonl"
    manifest = tmp_path / "manifest.json"
    result = run_pipeline(raw, out, manifest)
    assert out.exists()
    assert manifest.exists()
    assert result.total_records == 0
    manifest_obj = json.loads(manifest.read_text(encoding="utf-8"))
    assert manifest_obj["statistics"]["total_files"] == 0


def test_format_sheet_drops_empty_columns():
    rows = [
        ["a", "", "b"],
        ["1", "", "2"],
    ]
    rendered = _format_sheet_as_table(rows)
    assert "a" in rendered and "b" in rendered
    assert "1" in rendered and "2" in rendered
    # The empty middle column should have been dropped entirely.
    header_line = rendered.splitlines()[0]
    cells = [c.strip() for c in header_line.strip("|").split("|")]
    assert cells == ["a", "b"]
